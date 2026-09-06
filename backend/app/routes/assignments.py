from flask import Blueprint, request, jsonify, current_app
from flask_jwt_extended import jwt_required, get_jwt_identity
from datetime import datetime
import os
import uuid
import re
from werkzeug.utils import secure_filename
from app import db
from app.models import User, Assignment, Submission, Question, Answer
import openpyxl
from io import BytesIO

bp = Blueprint('assignments', __name__, url_prefix='/api/assignments')

ALLOWED_EXTENSIONS = {'png', 'jpg', 'jpeg', 'gif', 'pdf', 'doc', 'docx', 'py'}

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def parse_text_questions(text):
    questions = []
    lines = text.strip().split('\n')
    
    type_mapping = {
        '选择题': 'single_choice',
        '单选题': 'single_choice',
        '单选': 'single_choice',
        '多选题': 'multiple_choice',
        '多选': 'multiple_choice',
        '判断题': 'true_false',
        '判断': 'true_false',
        '填空题': 'fill_blank',
        '填空': 'fill_blank',
        '大题': 'text',
        '综合题': 'text',
        '综合大题': 'text',
        '简答题': 'text',
        '简答': 'text',
        '问答': 'text',
        '问答题': 'text',
        '论述题': 'text',
        '编程题': 'text',
        '编程': 'text',
    }
    
    type_pattern = r'^#+\s*[一二三四五六七八九十]+[、.．\s]\s*(.+?题|.+?选择|.+?判断|.+?填空|.+?大题|.+?编程|.+?问答|.+?简答)'
    question_pattern = r'^(\d+)[.、．]\s*(.+)$'
    single_option_pattern = r'^([A-D])[.、．)]\s*(.+)$'
    inline_option_pattern = r'([A-D])[.、．)]\s*(.+?)(?=\s{2,}[A-D][.、．)]|$)'
    answer_section_pattern = r'^#+\s*参考答案|^参考答案'
    answer_type_pattern = r'^#+\s*[一二三四五六七八九十]*[、.．\s]*\s*(单选|多选|判断|填空|问答|简答|编程|大题|综合|编程题)'
    sub_title_pattern = r'^###\s*(题目|题目要求|示例|要求)'
    
    current_type = 'single_choice'
    current_question = None
    in_answer_section = False
    current_answer_type = ''
    answers = {}
    big_question_counter = 0
    in_big_question = False
    in_code_block = False
    fill_blank_counter = {}
    text_counter = {}
    
    for i, line in enumerate(lines):
        original_line = line
        line = line.strip()
        
        if line.startswith('```'):
            in_code_block = not in_code_block
            if in_answer_section and current_answer_type == 'text':
                text_counter[current_answer_type] = text_counter.get(current_answer_type, 0)
                q_num = text_counter.get(current_answer_type, 0) + 9
                key = f"text_{q_num}"
                if key in answers:
                    answers[key] += '\n' + original_line
                else:
                    answers[key] = original_line
            continue
        
        if in_code_block:
            if in_answer_section and current_answer_type == 'text':
                text_counter[current_answer_type] = text_counter.get(current_answer_type, 0)
                q_num = text_counter.get(current_answer_type, 0) + 9
                key = f"text_{q_num}"
                if key in answers:
                    answers[key] += '\n' + line
                else:
                    answers[key] = line
            continue
        
        if not line or line == '---':
            continue
        
        if re.match(answer_section_pattern, line, re.IGNORECASE):
            in_answer_section = True
            if current_question:
                questions.append(current_question)
                current_question = None
            continue
        
        if in_answer_section:
            type_match = re.match(answer_type_pattern, line, re.IGNORECASE)
            if type_match:
                type_name = type_match.group(1).strip()
                matched = False
                for key, value in type_mapping.items():
                    if key in type_name or type_name in key:
                        current_answer_type = value
                        matched = True
                        break
                if not matched:
                    if '编程' in type_name:
                        current_answer_type = 'text'
                continue
            
            if current_answer_type in ['single_choice', 'multiple_choice']:
                pattern_with_num = r'(\d+)[.、．]\s*([A-D]+)'
                matches_with_num = re.findall(pattern_with_num, line)
                
                if matches_with_num:
                    for q_num, ans in matches_with_num:
                        key = f"{current_answer_type}_{q_num}"
                        answers[key] = ans.upper()
                else:
                    letter_matches = re.findall(r'([A-D]+)', line, re.IGNORECASE)
                    for idx, ans in enumerate(letter_matches, 1):
                        key = f"{current_answer_type}_{idx}"
                        answers[key] = ans.upper()
            
            elif current_answer_type == 'true_false':
                pattern_with_num = r'(\d+)[.、．]\s*([√×对错TFtf])'
                matches_with_num = re.findall(pattern_with_num, line)
                
                if matches_with_num:
                    for q_num, ans in matches_with_num:
                        if ans in ['√', '对', 'T', 't']:
                            ans_value = 'true'
                        else:
                            ans_value = 'false'
                        key = f"true_false_{q_num}"
                        answers[key] = ans_value
                else:
                    ans_list = re.findall(r'[√×对错TFtf]', line)
                    for idx, ans in enumerate(ans_list, 1):
                        if ans in ['√', '对', 'T', 't']:
                            ans_value = 'true'
                        else:
                            ans_value = 'false'
                        key = f"true_false_{idx}"
                        answers[key] = ans_value
            
            elif current_answer_type == 'fill_blank':
                pattern_with_num = r'^(\d+)[.、．]\s*(.+)$'
                match_with_num = re.match(pattern_with_num, line)
                
                if match_with_num:
                    q_num = match_with_num.group(1)
                    ans_content = match_with_num.group(2).strip()
                    key = f"fill_blank_{q_num}"
                    answers[key] = ans_content
                else:
                    fill_blank_counter[current_answer_type] = fill_blank_counter.get(current_answer_type, 0) + 1
                    q_num = fill_blank_counter[current_answer_type]
                    key = f"fill_blank_{q_num}"
                    answers[key] = line.strip()
            
            elif current_answer_type == 'text':
                if line.startswith('#') or line.startswith('---'):
                    continue
                text_counter[current_answer_type] = text_counter.get(current_answer_type, 0)
                text_counter[current_answer_type] += 1
                q_num = text_counter[current_answer_type] + 8
                key = f"text_{q_num}"
                if key in answers:
                    answers[key] += '\n' + line.strip()
                else:
                    answers[key] = line.strip()
            
            continue
        
        type_match = re.match(type_pattern, line)
        if type_match:
            if current_question:
                questions.append(current_question)
                current_question = None
            
            type_name = type_match.group(1).strip()
            for key, value in type_mapping.items():
                if key in type_name or type_name in key:
                    current_type = value
                    break
            
            if current_type == 'text':
                in_big_question = True
                big_question_counter += 1
                current_question = {
                    'question_type': 'text',
                    'content': '',
                    'score': 10,
                    'correct_answer': '',
                    'options': {'A': '', 'B': '', 'C': '', 'D': ''},
                    '_num': str(big_question_counter + 8)
                }
            else:
                in_big_question = False
            continue
        
        if in_big_question and current_question:
            if re.match(sub_title_pattern, line, re.IGNORECASE):
                continue
            
            if current_question['content']:
                current_question['content'] += '\n' + line
            else:
                current_question['content'] = line
            continue
        
        q_match = re.match(question_pattern, line)
        if q_match and not in_big_question:
            if current_question:
                questions.append(current_question)
            
            content = q_match.group(2).strip()
            content = re.sub(r'（\s*）|\(\s*\)|【\s*】', '', content)
            content = content.strip()
            
            default_scores = {
                'single_choice': 2,
                'true_false': 2,
                'multiple_choice': 3,
                'fill_blank': 5,
                'text': 10
            }
            
            current_question = {
                'question_type': current_type,
                'content': content,
                'score': default_scores.get(current_type, 10),
                'correct_answer': '',
                'options': {'A': '', 'B': '', 'C': '', 'D': ''},
                '_num': q_match.group(1)
            }
            continue
        
        if current_question and not in_big_question:
            inline_opts = re.findall(inline_option_pattern, line)
            if inline_opts and len(inline_opts) >= 1:
                for opt_letter, opt_content in inline_opts:
                    opt_content = opt_content.strip()
                    if opt_content:
                        current_question['options'][opt_letter] = opt_content
            else:
                single_opt_match = re.match(single_option_pattern, line)
                if single_opt_match:
                    opt_letter = single_opt_match.group(1)
                    opt_content = single_opt_match.group(2).strip()
                    current_question['options'][opt_letter] = opt_content
    
    if current_question:
        questions.append(current_question)
    
    for q in questions:
        q_num = q.get('_num', '')
        q_type = q['question_type']
        key = f"{q_type}_{q_num}"
        
        if key in answers:
            q['correct_answer'] = answers[key]
        
        if '_num' in q:
            del q['_num']
    
    return questions

@bp.route('', methods=['POST'])
@jwt_required()
def create_assignment():
    user_id = int(get_jwt_identity())
    user = User.query.get(user_id)
    
    if not user or user.role != 'teacher':
        return jsonify({'error': '只有老师可以创建作业'}), 403
    
    data = request.get_json()
    
    if not data or not data.get('title') or not data.get('due_date'):
        return jsonify({'error': '缺少必要字段'}), 400
    
    try:
        due_date = datetime.fromisoformat(data['due_date'].replace('Z', '+00:00'))
    except:
        return jsonify({'error': '日期格式错误'}), 400
    
    assignment = Assignment(
        title=data['title'],
        description=data.get('description', ''),
        due_date=due_date,
        teacher_id=user_id,
        course_id=data.get('course_id'),
        total_score=data.get('total_score', 100)
    )
    
    db.session.add(assignment)
    db.session.commit()
    
    return jsonify({
        'message': '作业创建成功',
        'assignment': assignment.to_dict()
    }), 201

@bp.route('', methods=['GET'])
@jwt_required()
def get_assignments():
    user_id = int(get_jwt_identity())
    user = User.query.get(user_id)

    if user.role == 'teacher':
        query = Assignment.query.filter_by(teacher_id=user_id)
    else:
        query = Assignment.query

    # 支持分页（传 page 参数启用），不传则全量返回以兼容旧前端
    page = request.args.get('page', type=int)
    if page:
        per_page = request.args.get('per_page', 20, type=int)
        paginated = query.order_by(Assignment.id.desc()).paginate(
            page=page, per_page=per_page, error_out=False
        )
        return jsonify({
            'assignments': [a.to_dict() for a in paginated.items],
            'total': paginated.total,
            'page': page,
            'per_page': per_page,
            'has_next': paginated.has_next
        }), 200

    assignments = query.all()
    return jsonify({
        'assignments': [assignment.to_dict() for assignment in assignments]
    }), 200

@bp.route('/<int:assignment_id>', methods=['GET'])
@jwt_required()
def get_assignment(assignment_id):
    user_id = int(get_jwt_identity())
    user = User.query.get(user_id)
    
    assignment = Assignment.query.get(assignment_id)
    
    if not assignment:
        return jsonify({'error': '作业不存在'}), 404
    
    assignment_dict = assignment.to_dict()
    
    questions = Question.query.filter_by(assignment_id=assignment_id).order_by(Question.question_number).all()
    assignment_dict['questions'] = [q.to_dict() for q in questions]
    
    if user.role == 'student':
        submission = Submission.query.filter_by(
            assignment_id=assignment_id,
            student_id=user_id
        ).first()
        if submission:
            assignment_dict['my_submission'] = submission.to_dict()
            answers = Answer.query.filter_by(submission_id=submission.id).all()
            assignment_dict['my_answers'] = {a.question_id: a.to_dict() for a in answers}
    
    return jsonify({'assignment': assignment_dict}), 200

@bp.route('/<int:assignment_id>', methods=['PUT'])
@jwt_required()
def update_assignment(assignment_id):
    user_id = int(get_jwt_identity())
    user = User.query.get(user_id)
    
    if not user or user.role != 'teacher':
        return jsonify({'error': '只有老师可以修改作业'}), 403
    
    assignment = Assignment.query.get(assignment_id)
    
    if not assignment:
        return jsonify({'error': '作业不存在'}), 404
    
    if assignment.teacher_id != user_id:
        return jsonify({'error': '无权操作此作业'}), 403
    
    data = request.get_json()
    
    if data.get('title'):
        assignment.title = data['title']
    if data.get('description') is not None:
        assignment.description = data['description']
    if data.get('due_date'):
        try:
            assignment.due_date = datetime.fromisoformat(data['due_date'].replace('Z', '+00:00'))
        except:
            pass
    if data.get('total_score') is not None:
        assignment.total_score = data['total_score']
    
    db.session.commit()
    
    return jsonify({
        'message': '作业更新成功',
        'assignment': assignment.to_dict()
    }), 200

@bp.route('/<int:assignment_id>/submit', methods=['POST'])
@jwt_required()
def submit_assignment(assignment_id):
    user_id = int(get_jwt_identity())
    user = User.query.get(user_id)
    
    if not user or user.role != 'student':
        return jsonify({'error': '只有学生可以提交作业'}), 403
    
    assignment = Assignment.query.get(assignment_id)
    
    if not assignment:
        return jsonify({'error': '作业不存在'}), 404
    
    if assignment.due_date < datetime.utcnow():
        return jsonify({'error': '作业已截止，无法提交'}), 400
    
    existing_submission = Submission.query.filter_by(
        assignment_id=assignment_id,
        student_id=user_id
    ).first()
    
    if request.content_type and 'multipart/form-data' in request.content_type:
        content = request.form.get('content', '')
        file = request.files.get('file')
        answers_data = request.form.get('answers')
        
        file_url = existing_submission.file_url if existing_submission else None
        if file and allowed_file(file.filename):
            filename = secure_filename(file.filename)
            unique_filename = f"{uuid.uuid4().hex}_{filename}"
            
            upload_folder = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'tupian')
            os.makedirs(upload_folder, exist_ok=True)
            
            file_path = os.path.join(upload_folder, unique_filename)
            file.save(file_path)
            
            file_url = f"/tupian/{unique_filename}"
    else:
        data = request.get_json()
        content = data.get('content', '')
        file_url = data.get('file_url')
        answers_data = data.get('answers')
    
    if existing_submission:
        existing_submission.content = content
        if file_url:
            existing_submission.file_url = file_url
        existing_submission.submitted_at = datetime.utcnow()
        existing_submission.status = 'submitted'
        submission = existing_submission
    else:
        submission = Submission(
            assignment_id=assignment_id,
            student_id=user_id,
            content=content,
            file_url=file_url
        )
        db.session.add(submission)
        db.session.flush()
    
    if answers_data:
        import json
        try:
            if isinstance(answers_data, str):
                answers_list = json.loads(answers_data)
            else:
                answers_list = answers_data
            
            for answer_item in answers_list:
                question_id = answer_item.get('question_id')
                if not question_id:
                    continue
                
                question = Question.query.get(question_id)
                if not question or question.assignment_id != assignment_id:
                    continue
                
                existing_answer = Answer.query.filter_by(
                    submission_id=submission.id,
                    question_id=question_id
                ).first()
                
                if existing_answer:
                    existing_answer.answer_text = answer_item.get('answer_text')
                    existing_answer.answer_image_url = answer_item.get('answer_image_url')
                    existing_answer.updated_at = datetime.utcnow()
                else:
                    answer = Answer(
                        submission_id=submission.id,
                        question_id=question_id,
                        answer_text=answer_item.get('answer_text'),
                        answer_image_url=answer_item.get('answer_image_url')
                    )
                    db.session.add(answer)
        except Exception as e:
            print(f"处理答案数据错误: {e}")
    
    db.session.commit()
    
    return jsonify({
        'message': '作业提交成功',
        'submission': submission.to_dict()
    }), 200

@bp.route('/<int:assignment_id>/submissions', methods=['GET'])
@jwt_required()
def get_submissions(assignment_id):
    user_id = int(get_jwt_identity())
    user = User.query.get(user_id)
    
    if not user or user.role != 'teacher':
        return jsonify({'error': '只有老师可以查看提交'}), 403
    
    assignment = Assignment.query.get(assignment_id)
    
    if not assignment:
        return jsonify({'error': '作业不存在'}), 404
    
    if assignment.teacher_id != user_id:
        return jsonify({'error': '无权查看此作业的提交'}), 403
    
    questions = Question.query.filter_by(assignment_id=assignment_id).order_by(Question.question_number).all()
    questions_data = [q.to_dict() for q in questions]
    
    submissions = Submission.query.filter_by(assignment_id=assignment_id).all()
    
    submissions_data = []
    for submission in submissions:
        submission_dict = submission.to_dict()
        student = User.query.get(submission.student_id)
        if student:
            submission_dict['student_name'] = student.name
        
        answers = Answer.query.filter_by(submission_id=submission.id).all()
        submission_dict['answers'] = {a.question_id: a.to_dict() for a in answers}
        
        submissions_data.append(submission_dict)
    
    return jsonify({
        'submissions': submissions_data,
        'questions': questions_data
    }), 200

@bp.route('/submissions/<int:submission_id>/grade', methods=['PUT'])
@jwt_required()
def grade_submission(submission_id):
    user_id = int(get_jwt_identity())
    user = User.query.get(user_id)
    
    if not user or user.role != 'teacher':
        return jsonify({'error': '只有老师可以批改作业'}), 403
    
    submission = Submission.query.get(submission_id)
    
    if not submission:
        return jsonify({'error': '提交不存在'}), 404
    
    assignment = Assignment.query.get(submission.assignment_id)
    
    if assignment.teacher_id != user_id:
        return jsonify({'error': '无权批改此作业'}), 403
    
    data = request.get_json()
    
    auto_grade = data.get('auto_grade', False)
    
    if auto_grade:
        questions = Question.query.filter_by(assignment_id=assignment.id).all()
        answers = Answer.query.filter_by(submission_id=submission_id).all()
        answers_dict = {a.question_id: a for a in answers}
        
        auto_score = 0
        auto_graded_answers = []
        
        for question in questions:
            answer = answers_dict.get(question.id)
            if not answer:
                continue
            
            if question.question_type == 'single_choice':
                if answer.answer_text == question.correct_answer:
                    answer.score = question.score
                    answer.is_correct = True
                    auto_score += question.score
                else:
                    answer.score = 0
                    answer.is_correct = False
                auto_graded_answers.append({
                    'question_id': question.id,
                    'score': answer.score,
                    'is_correct': answer.is_correct
                })
            
            elif question.question_type == 'multiple_choice':
                correct_set = set(question.correct_answer.split(',')) if question.correct_answer else set()
                student_set = set(answer.answer_text.split(',')) if answer.answer_text else set()
                
                if not student_set:
                    answer.score = 0
                    answer.is_correct = False
                elif student_set == correct_set:
                    answer.score = question.score
                    answer.is_correct = True
                    auto_score += question.score
                elif student_set.issubset(correct_set):
                    answer.score = question.score / 2
                    answer.is_correct = False
                    auto_score += question.score / 2
                else:
                    answer.score = 0
                    answer.is_correct = False
                auto_graded_answers.append({
                    'question_id': question.id,
                    'score': answer.score,
                    'is_correct': answer.is_correct
                })
            
            elif question.question_type == 'fill_blank':
                correct_answers = [ans.strip().lower() for ans in question.correct_answer.split('|')] if question.correct_answer else []
                student_answer = answer.answer_text.strip().lower() if answer.answer_text else ''
                
                if not student_answer:
                    answer.score = 0
                    answer.is_correct = False
                elif student_answer in correct_answers:
                    answer.score = question.score
                    answer.is_correct = True
                    auto_score += question.score
                else:
                    answer.score = 0
                    answer.is_correct = False
                auto_graded_answers.append({
                    'question_id': question.id,
                    'score': answer.score,
                    'is_correct': answer.is_correct
                })
            
            elif question.question_type == 'true_false':
                correct_answer = question.correct_answer if question.correct_answer else ''
                student_answer = answer.answer_text if answer.answer_text else ''
                
                if not student_answer:
                    answer.score = 0
                    answer.is_correct = False
                elif student_answer == correct_answer:
                    answer.score = question.score
                    answer.is_correct = True
                    auto_score += question.score
                else:
                    answer.score = 0
                    answer.is_correct = False
                auto_graded_answers.append({
                    'question_id': question.id,
                    'score': answer.score,
                    'is_correct': answer.is_correct
                })
        
        submission.score = auto_score
        
        db.session.commit()
        
        return jsonify({
            'message': '自动批改完成',
            'submission': submission.to_dict(),
            'auto_score': auto_score,
            'answers': auto_graded_answers
        }), 200
    
    if data.get('score') is not None:
        submission.score = data['score']
    if data.get('feedback') is not None:
        submission.feedback = data['feedback']
    if data.get('overall_comment') is not None:
        submission.overall_comment = data['overall_comment']
    submission.status = 'graded'
    submission.graded_at = datetime.utcnow()
    
    if data.get('answers'):
        for answer_data in data['answers']:
            answer_id = answer_data.get('answer_id')
            question_id = answer_data.get('question_id')
            
            if answer_id:
                answer = Answer.query.get(answer_id)
            elif question_id:
                answer = Answer.query.filter_by(
                    submission_id=submission_id,
                    question_id=question_id
                ).first()
            else:
                continue
            
            if answer:
                if answer_data.get('score') is not None:
                    answer.score = answer_data['score']
                if answer_data.get('is_correct') is not None:
                    answer.is_correct = answer_data['is_correct']
                if answer_data.get('feedback') is not None:
                    answer.feedback = answer_data['feedback']
    
    db.session.commit()
    
    return jsonify({
        'message': '批改成功',
        'submission': submission.to_dict()
    }), 200

@bp.route('/my-submissions', methods=['GET'])
@jwt_required()
def get_my_submissions():
    user_id = int(get_jwt_identity())
    user = User.query.get(user_id)
    
    if not user or user.role != 'student':
        return jsonify({'error': '只有学生可以查看自己的提交'}), 403
    
    submissions = Submission.query.filter_by(student_id=user_id).all()
    
    submissions_data = []
    for submission in submissions:
        submission_dict = submission.to_dict()
        answers = Answer.query.filter_by(submission_id=submission.id).all()
        submission_dict['answers'] = {a.question_id: a.to_dict() for a in answers}
        submissions_data.append(submission_dict)
    
    return jsonify({
        'submissions': submissions_data
    }), 200

@bp.route('/ai-grade', methods=['POST'])
@jwt_required()
def ai_grade_question():
    user_id = int(get_jwt_identity())
    user = User.query.get(user_id)
    
    if not user or user.role != 'teacher':
        return jsonify({'error': '只有老师可以使用AI评分功能'}), 403
    
    data = request.get_json()
    
    question_id = data.get('question_id')
    submission_id = data.get('submission_id')
    
    if not question_id or not submission_id:
        return jsonify({'error': '缺少题目ID或提交ID'}), 400
    
    question = Question.query.get(question_id)
    if not question:
        return jsonify({'error': '题目不存在'}), 404
    
    if question.question_type != 'text':
        return jsonify({'error': 'AI评分仅支持非选择题'}), 400
    
    submission = Submission.query.get(submission_id)
    if not submission:
        return jsonify({'error': '提交不存在'}), 404
    
    assignment = Assignment.query.get(question.assignment_id)
    if assignment.teacher_id != user_id:
        return jsonify({'error': '无权批改此作业'}), 403
    
    answer = Answer.query.filter_by(
        submission_id=submission_id,
        question_id=question_id
    ).first()
    
    if not answer:
        return jsonify({'error': '未找到学生答案'}), 404
    
    from app.services import DeepSeekService
    
    result = DeepSeekService.grade_text_question(
        question_content=question.content,
        question_score=question.score,
        reference_answer=question.correct_answer,
        student_answer=answer.answer_text
    )
    
    if result['success']:
        answer.score = result['score']
        answer.feedback = result.get('feedback', '')
        answer.is_correct = result['score'] == question.score
        db.session.commit()
        
        return jsonify({
            'success': True,
            'score': result['score'],
            'feedback': result.get('feedback', ''),
            'analysis': result.get('analysis', ''),
            'answer': answer.to_dict()
        }), 200
    else:
        error_msg = result.get('error', 'AI评分失败')
        if '402' in error_msg or 'Insufficient Balance' in error_msg:
            error_msg = 'DeepSeek账户余额不足，请前往 https://platform.deepseek.com/ 充值后再使用AI评分功能'
        return jsonify({
            'success': False,
            'error': error_msg
        }), 200

@bp.route('/submissions/<int:submission_id>/generate-comment', methods=['POST'])
@jwt_required()
def generate_overall_comment(submission_id):
    user_id = int(get_jwt_identity())
    user = User.query.get(user_id)
    
    if not user or user.role != 'teacher':
        return jsonify({'error': '只有老师可以使用AI生成总评功能'}), 403
    
    submission = Submission.query.get(submission_id)
    
    if not submission:
        return jsonify({'error': '提交不存在'}), 404
    
    assignment = Assignment.query.get(submission.assignment_id)
    
    if assignment.teacher_id != user_id:
        return jsonify({'error': '无权批改此作业'}), 403
    
    if submission.comment_locked:
        return jsonify({'error': '评语已锁定，无法重新生成'}), 400
    
    questions = Question.query.filter_by(assignment_id=assignment.id).all()
    answers = Answer.query.filter_by(submission_id=submission_id).all()
    answers_dict = {a.question_id: a for a in answers}
    
    question_results = []
    for q in questions:
        answer = answers_dict.get(q.id)
        type_names = {
            'single_choice': '单选题',
            'multiple_choice': '多选题',
            'fill_blank': '填空题',
            'true_false': '判断题',
            'text': '大题'
        }
        question_results.append({
            'type': type_names.get(q.question_type, '未知题型'),
            'score': answer.score if answer else 0,
            'max_score': q.score,
            'feedback': answer.feedback if answer else ''
        })
    
    student = User.query.get(submission.student_id)
    student_name = student.name if student else None
    
    from app.services import DeepSeekService
    
    result = DeepSeekService.generate_overall_comment(
        assignment_title=assignment.title,
        student_name=student_name,
        total_score=submission.score or 0,
        max_score=assignment.total_score or 100,
        question_results=question_results
    )
    
    if result['success']:
        submission.overall_comment = result['comment']
        db.session.commit()
        
        return jsonify({
            'success': True,
            'comment': result['comment'],
            'submission': submission.to_dict()
        }), 200
    else:
        error_msg = result.get('error', 'AI生成总评失败')
        if '402' in error_msg or 'Insufficient Balance' in error_msg:
            error_msg = 'DeepSeek账户余额不足，请前往 https://platform.deepseek.com/ 充值后再使用AI评分功能'
        return jsonify({
            'success': False,
            'error': error_msg
        }), 200

@bp.route('/submissions/<int:submission_id>/lock-comment', methods=['POST'])
@jwt_required()
def lock_comment(submission_id):
    user_id = int(get_jwt_identity())
    user = User.query.get(user_id)
    
    if not user or user.role != 'teacher':
        return jsonify({'error': '只有老师可以锁定评语'}), 403
    
    submission = Submission.query.get(submission_id)
    
    if not submission:
        return jsonify({'error': '提交不存在'}), 404
    
    assignment = Assignment.query.get(submission.assignment_id)
    
    if assignment.teacher_id != user_id:
        return jsonify({'error': '无权操作此作业'}), 403
    
    data = request.get_json() or {}
    action = data.get('action', 'lock')
    
    if action == 'lock':
        submission.comment_locked = True
        message = '评语已锁定'
    else:
        submission.comment_locked = False
        message = '评语已解锁'
    
    db.session.commit()
    
    return jsonify({
        'success': True,
        'message': message,
        'comment_locked': submission.comment_locked,
        'submission': submission.to_dict()
    }), 200

@bp.route('/submissions/<int:submission_id>/update-comment', methods=['PUT'])
@jwt_required()
def update_comment(submission_id):
    user_id = int(get_jwt_identity())
    user = User.query.get(user_id)
    
    if not user or user.role != 'teacher':
        return jsonify({'error': '只有老师可以修改评语'}), 403
    
    submission = Submission.query.get(submission_id)
    
    if not submission:
        return jsonify({'error': '提交不存在'}), 404
    
    assignment = Assignment.query.get(submission.assignment_id)
    
    if assignment.teacher_id != user_id:
        return jsonify({'error': '无权操作此作业'}), 403
    
    if submission.comment_locked:
        return jsonify({'error': '评语已锁定，无法修改'}), 400
    
    data = request.get_json()
    
    if data.get('overall_comment') is not None:
        submission.overall_comment = data['overall_comment']
    
    if data.get('score') is not None:
        submission.score = data['score']
    
    if data.get('feedback') is not None:
        submission.feedback = data['feedback']
    
    db.session.commit()
    
    return jsonify({
        'success': True,
        'message': '评语更新成功',
        'submission': submission.to_dict()
    }), 200

@bp.route('/ai-grade-python', methods=['POST'])
@jwt_required()
def ai_grade_python_code():
    user_id = int(get_jwt_identity())
    user = User.query.get(user_id)
    
    if not user or user.role != 'teacher':
        return jsonify({'error': '只有老师可以使用AI评分功能'}), 403
    
    data = request.get_json()
    
    question_id = data.get('question_id')
    submission_id = data.get('submission_id')
    
    if not question_id or not submission_id:
        return jsonify({'error': '缺少题目ID或提交ID'}), 400
    
    question = Question.query.get(question_id)
    if not question:
        return jsonify({'error': '题目不存在'}), 404
    
    submission = Submission.query.get(submission_id)
    if not submission:
        return jsonify({'error': '提交不存在'}), 404
    
    assignment = Assignment.query.get(question.assignment_id)
    if assignment.teacher_id != user_id:
        return jsonify({'error': '无权批改此作业'}), 403
    
    answer = Answer.query.filter_by(
        submission_id=submission_id,
        question_id=question_id
    ).first()
    
    if not answer:
        return jsonify({'error': '未找到学生答案'}), 404
    
    student_code = answer.answer_text
    
    if not student_code:
        if submission.file_url and submission.file_url.endswith('.py'):
            try:
                file_path = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'tupian', submission.file_url.split('/')[-1])
                if os.path.exists(file_path):
                    with open(file_path, 'r', encoding='utf-8') as f:
                        student_code = f.read()
            except Exception as e:
                return jsonify({
                    'success': False,
                    'error': f'读取Python文件失败: {str(e)}'
                }), 200
    
    if not student_code:
        return jsonify({
            'success': False,
            'error': '学生未提交代码'
        }), 200
    
    from app.services import DeepSeekService
    
    result = DeepSeekService.grade_python_code(
        question_content=question.content,
        question_score=question.score,
        reference_answer=question.correct_answer,
        student_code=student_code,
        requirements=None
    )
    
    if result['success']:
        answer.score = result['score']
        answer.feedback = result.get('feedback', '')
        answer.is_correct = result['score'] == question.score
        db.session.commit()
        
        return jsonify({
            'success': True,
            'score': result['score'],
            'feedback': result.get('feedback', ''),
            'code_analysis': result.get('code_analysis', {}),
            'improved_code': result.get('improved_code', ''),
            'answer': answer.to_dict()
        }), 200
    else:
        error_msg = result.get('error', 'AI评分失败')
        if '402' in error_msg or 'Insufficient Balance' in error_msg:
            error_msg = 'DeepSeek账户余额不足，请前往 https://platform.deepseek.com/ 充值后再使用AI评分功能'
        return jsonify({
            'success': False,
            'error': error_msg
        }), 200

@bp.route('/read-python-file', methods=['POST'])
@jwt_required()
def read_python_file():
    user_id = int(get_jwt_identity())
    user = User.query.get(user_id)
    
    data = request.get_json()
    file_url = data.get('file_url')
    
    if not file_url:
        return jsonify({'error': '缺少文件URL'}), 400
    
    if not file_url.endswith('.py'):
        return jsonify({'error': '只支持Python文件'}), 400
    
    try:
        file_path = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'tupian', file_url.split('/')[-1])
        if not os.path.exists(file_path):
            return jsonify({'error': '文件不存在'}), 404
        
        with open(file_path, 'r', encoding='utf-8') as f:
            code_content = f.read()
        
        return jsonify({
            'success': True,
            'code': code_content
        }), 200
    except Exception as e:
        return jsonify({
            'success': False,
            'error': f'读取文件失败: {str(e)}'
        }), 500

@bp.route('/<int:assignment_id>/import-questions', methods=['POST'])
@jwt_required()
def import_questions(assignment_id):
    user_id = int(get_jwt_identity())
    user = User.query.get(user_id)
    
    if not user or user.role != 'teacher':
        return jsonify({'error': '只有老师可以导入题目'}), 403
    
    assignment = Assignment.query.get(assignment_id)
    
    if not assignment:
        return jsonify({'error': '作业不存在'}), 404
    
    if assignment.teacher_id != user_id:
        return jsonify({'error': '无权操作此作业'}), 403
    
    if 'file' not in request.files:
        return jsonify({'error': '没有上传文件'}), 400
    
    file = request.files['file']
    
    if not file.filename:
        return jsonify({'error': '没有选择文件'}), 400
    
    if not file.filename.endswith(('.xlsx', '.xls')):
        return jsonify({'error': '只支持Excel文件（.xlsx或.xls）'}), 400
    
    try:
        wb = openpyxl.load_workbook(BytesIO(file.read()))
        ws = wb.active
        
        headers = [str(cell.value).strip().lower() if cell.value else '' for cell in ws[1]]
        
        required_headers = ['题型', '题目内容', '分数']
        for header in required_headers:
            if header not in headers:
                return jsonify({'error': f'Excel缺少必需列：{header}'}), 400
        
        type_map = {
            '单选题': 'single_choice',
            '多选题': 'multiple_choice',
            '填空题': 'fill_blank',
            '判断题': 'true_false',
            '大题': 'text',
            '简答题': 'text',
            '主观题': 'text'
        }
        
        success_count = 0
        error_list = []
        questions_data = []
        
        existing_count = Question.query.filter_by(assignment_id=assignment_id).count()
        question_number = existing_count + 1
        
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not any(row):
                continue
            
            row_data = dict(zip(headers, row))
            
            question_type_str = str(row_data.get('题型', '')).strip()
            question_type = type_map.get(question_type_str)
            
            if not question_type:
                error_list.append(f'第{row_idx}行：题型"{question_type_str}"不支持，支持：单选题、多选题、填空题、判断题、大题')
                continue
            
            content = str(row_data.get('题目内容', '')).strip()
            if not content:
                error_list.append(f'第{row_idx}行：题目内容不能为空')
                continue
            
            try:
                score = float(row_data.get('分数', 0))
                if score <= 0:
                    error_list.append(f'第{row_idx}行：分数必须大于0')
                    continue
            except:
                error_list.append(f'第{row_idx}行：分数格式错误')
                continue
            
            correct_answer = str(row_data.get('正确答案', '')).strip() if row_data.get('正确答案') else ''
            
            options = {'A': '', 'B': '', 'C': '', 'D': ''}
            if question_type in ['single_choice', 'multiple_choice']:
                for opt in ['A', 'B', 'C', 'D']:
                    opt_key = f'选项{opt}'
                    if opt_key in headers:
                        opt_val = row_data.get(opt_key)
                        options[opt] = str(opt_val).strip() if opt_val else ''
                
                if not correct_answer:
                    error_list.append(f'第{row_idx}行：选择题必须填写正确答案')
                    continue
            
            if question_type == 'true_false':
                if correct_answer not in ['正确', '错误', 'true', 'false', 'True', 'False', '对', '错']:
                    error_list.append(f'第{row_idx}行：判断题正确答案必须是"正确"或"错误"')
                    continue
                correct_answer = 'true' if correct_answer in ['正确', 'true', 'True', '对'] else 'false'
            
            question = Question(
                assignment_id=assignment_id,
                question_number=question_number,
                question_type=question_type,
                content=content,
                score=score,
                correct_answer=correct_answer
            )
            
            if options and any(options.values()):
                question.set_options(options)
            
            db.session.add(question)
            question_number += 1
            success_count += 1
            questions_data.append({
                'question_type': question_type,
                'content': content,
                'score': score,
                'correct_answer': correct_answer,
                'options': options
            })
        
        db.session.commit()
        
        return jsonify({
            'success': True,
            'message': f'成功导入 {success_count} 道题目',
            'success_count': success_count,
            'error_count': len(error_list),
            'errors': error_list[:10],
            'questions': questions_data
        }), 200
        
    except Exception as e:
        db.session.rollback()
        return jsonify({'error': f'导入失败：{str(e)}'}), 500

@bp.route('/ai-status', methods=['GET'])
@jwt_required()
def check_ai_status():
    user_id = int(get_jwt_identity())
    user = User.query.get(user_id)
    
    if not user or user.role != 'teacher':
        return jsonify({'error': '只有老师可以检查AI状态'}), 403
    
    from app.services.ai_parser import AIQuestionParser
    status = AIQuestionParser.check_status()
    
    if status['available']:
        return jsonify({
            'status': 'ok',
            'configured': True,
            'message': status['message']
        }), 200
    else:
        return jsonify({
            'status': 'error',
            'configured': False,
            'message': status['message'],
            'hint': status.get('hint', '')
        }), 200

@bp.route('/ai-test', methods=['POST'])
@jwt_required()
def test_ai_connection():
    user_id = int(get_jwt_identity())
    user = User.query.get(user_id)
    
    if not user or user.role != 'teacher':
        return jsonify({'error': '只有老师可以测试AI连接'}), 403
    
    try:
        from app.services.ai_parser import AIQuestionParser
        result = AIQuestionParser.test_connection()
        
        if result['success']:
            return jsonify({
                'success': True,
                'message': result['message']
            }), 200
        else:
            return jsonify({
                'success': False,
                'error': result['error'],
                'hint': result.get('hint', '')
            }), 200
    except Exception as e:
        print(f"[AI Test Route] Error: {e}")
        import traceback
        traceback.print_exc()
        return jsonify({'error': f'测试失败: {str(e)}'}), 500

@bp.route('/<int:assignment_id>/ai-parse', methods=['POST'])
@jwt_required()
def ai_parse_questions(assignment_id):
    user_id = int(get_jwt_identity())
    user = User.query.get(user_id)
    
    if not user or user.role != 'teacher':
        return jsonify({'error': '只有老师可以使用AI解析功能'}), 403
    
    assignment = Assignment.query.get(assignment_id)
    
    if not assignment:
        return jsonify({'error': '作业不存在'}), 404
    
    if assignment.teacher_id != user_id:
        return jsonify({'error': '无权操作此作业'}), 403
    
    data = request.get_json()
    text = data.get('text', '')
    
    if not text.strip():
        return jsonify({'error': '文本内容不能为空'}), 400
    
    try:
        from app.services.ai_parser import AIQuestionParser
        
        result = AIQuestionParser.parse_questions(text)
        
        if not result['success']:
            error_msg = result.get('error', 'AI解析失败')
            return jsonify({'error': error_msg}), 400
        
        questions = result['questions']
        
        if not questions:
            return jsonify({'error': 'AI未能识别出有效题目'}), 400
        
        success_count = 0
        questions_data = []
        
        existing_count = Question.query.filter_by(assignment_id=assignment_id).count()
        question_number = existing_count + 1
        
        for q in questions:
            question = Question(
                assignment_id=assignment_id,
                question_number=question_number,
                question_type=q['question_type'],
                content=q['content'],
                score=q['score'],
                correct_answer=q.get('correct_answer', '')
            )
            
            if q.get('options') and any(q['options'].values()):
                question.set_options(q['options'])
            
            db.session.add(question)
            question_number += 1
            success_count += 1
            questions_data.append({
                'question_type': q['question_type'],
                'content': q['content'],
                'score': q['score'],
                'correct_answer': q.get('correct_answer', ''),
                'options': q.get('options', {'A': '', 'B': '', 'C': '', 'D': ''})
            })
        
        db.session.commit()
        
        return jsonify({
            'success': True,
            'message': f'AI成功解析并导入 {success_count} 道题目',
            'success_count': success_count,
            'error_count': 0,
            'questions': questions_data
        }), 200
        
    except Exception as e:
        print(f"[AI Parse Route] Error: {e}")
        import traceback
        traceback.print_exc()
        db.session.rollback()
        return jsonify({'error': f'解析失败: {str(e)}'}), 500

@bp.route('/<int:assignment_id>/clear-questions', methods=['DELETE'])
@jwt_required()
def clear_assignment_questions(assignment_id):
    user_id = int(get_jwt_identity())
    user = User.query.get(user_id)
    
    if not user or user.role != 'teacher':
        return jsonify({'error': '只有老师可以清除题目'}), 403
    
    assignment = Assignment.query.get(assignment_id)
    
    if not assignment:
        return jsonify({'error': '作业不存在'}), 404
    
    if assignment.teacher_id != user_id:
        return jsonify({'error': '无权操作此作业'}), 403
    
    try:
        deleted_count = Question.query.filter_by(assignment_id=assignment_id).delete()
        assignment.total_score = 0
        db.session.commit()
        
        return jsonify({
            'success': True,
            'message': f'已清除 {deleted_count} 道题目',
            'deleted_count': deleted_count
        }), 200
    except Exception as e:
        db.session.rollback()
        return jsonify({'error': f'清除失败：{str(e)}'}), 500

@bp.route('/import-template', methods=['GET'])
@jwt_required()
def download_questions_template():
    user_id = int(get_jwt_identity())
    user = User.query.get(user_id)
    
    if not user or user.role != 'teacher':
        return jsonify({'error': '只有老师可以下载模板'}), 403
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '题目导入模板'
    
    headers = ['题型', '题目内容', '分数', '正确答案', '选项A', '选项B', '选项C', '选项D']
    ws.append(headers)
    
    example_data = [
        ['单选题', 'Python中用于输出的函数是？', 2, 'A', 'print()', 'input()', 'output()', 'display()'],
        ['多选题', '以下哪些是Python的数据类型？', 3, 'A,B,C', 'int', 'str', 'list', 'array'],
        ['填空题', 'Python中用于定义函数的关键字是____。', 5, 'def|DEF', '', '', '', ''],
        ['判断题', 'Python是一种解释型语言。', 2, '正确', '', '', '', ''],
        ['大题', '请简述Python的特点。', 10, '', '', '', '', '']
    ]
    
    for data in example_data:
        ws.append(data)
    
    ws2 = wb.create_sheet('题型说明')
    ws2.append(['题型', '说明', '正确答案格式'])
    ws2.append(['单选题', '从A、B、C、D中选择一个正确答案', '填写单个字母，如：A'])
    ws2.append(['多选题', '从A、B、C、D中选择多个正确答案', '多个字母用逗号分隔，如：A,B,C'])
    ws2.append(['填空题', '学生填写答案', '多个正确答案用|分隔，如：答案1|答案2'])
    ws2.append(['判断题', '判断对错', '填写"正确"或"错误"'])
    ws2.append(['大题', '主观题，需要老师批改', '可不填写正确答案'])
    
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    from flask import send_file
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name='题目导入模板.xlsx'
    )

@bp.route('/<int:assignment_id>/import-text', methods=['POST'])
@jwt_required()
def import_text_questions(assignment_id):
    user_id = int(get_jwt_identity())
    user = User.query.get(user_id)
    
    if not user or user.role != 'teacher':
        return jsonify({'error': '只有老师可以导入题目'}), 403
    
    assignment = Assignment.query.get(assignment_id)
    
    if not assignment:
        return jsonify({'error': '作业不存在'}), 404
    
    if assignment.teacher_id != user_id:
        return jsonify({'error': '无权操作此作业'}), 403
    
    data = request.get_json()
    text = data.get('text', '')
    
    if not text.strip():
        return jsonify({'error': '文本内容不能为空'}), 400
    
    try:
        questions = parse_text_questions(text)
        
        if not questions:
            return jsonify({'error': '未能识别出有效题目，请检查格式'}), 400
        
        success_count = 0
        questions_data = []
        
        existing_count = Question.query.filter_by(assignment_id=assignment_id).count()
        question_number = existing_count + 1
        
        for q in questions:
            question = Question(
                assignment_id=assignment_id,
                question_number=question_number,
                question_type=q['question_type'],
                content=q['content'],
                score=q['score'],
                correct_answer=q.get('correct_answer', '')
            )
            
            if q.get('options') and any(q['options'].values()):
                question.set_options(q['options'])
            
            db.session.add(question)
            question_number += 1
            success_count += 1
            questions_data.append({
                'question_type': q['question_type'],
                'content': q['content'],
                'score': q['score'],
                'correct_answer': q.get('correct_answer', ''),
                'options': q.get('options', {'A': '', 'B': '', 'C': '', 'D': ''})
            })
        
        db.session.commit()
        
        return jsonify({
            'success': True,
            'message': f'成功导入 {success_count} 道题目',
            'success_count': success_count,
            'error_count': 0,
            'questions': questions_data
        }), 200
        
    except Exception as e:
        db.session.rollback()
        return jsonify({'error': f'导入失败：{str(e)}'}), 500

@bp.route('/<int:assignment_id>/import-file', methods=['POST'])
@jwt_required()
def import_file_questions(assignment_id):
    user_id = int(get_jwt_identity())
    user = User.query.get(user_id)
    
    if not user or user.role != 'teacher':
        return jsonify({'error': '只有老师可以导入题目'}), 403
    
    assignment = Assignment.query.get(assignment_id)
    
    if not assignment:
        return jsonify({'error': '作业不存在'}), 404
    
    if assignment.teacher_id != user_id:
        return jsonify({'error': '无权操作此作业'}), 403
    
    if 'file' not in request.files:
        return jsonify({'error': '没有上传文件'}), 400
    
    file = request.files['file']
    
    if not file.filename:
        return jsonify({'error': '没有选择文件'}), 400
    
    if not file.filename.endswith('.txt'):
        return jsonify({'error': '只支持文本文件（.txt）'}), 400
    
    try:
        try:
            content = file.read().decode('utf-8')
        except UnicodeDecodeError:
            try:
                content = file.read().decode('gbk')
            except:
                return jsonify({'error': '文件编码不支持，请使用UTF-8或GBK编码'}), 400
        
        questions = parse_text_questions(content)
        
        if not questions:
            return jsonify({'error': '未能识别出有效题目，请检查格式'}), 400
        
        success_count = 0
        questions_data = []
        
        existing_count = Question.query.filter_by(assignment_id=assignment_id).count()
        question_number = existing_count + 1
        
        for q in questions:
            question = Question(
                assignment_id=assignment_id,
                question_number=question_number,
                question_type=q['question_type'],
                content=q['content'],
                score=q['score'],
                correct_answer=q.get('correct_answer', '')
            )
            
            if q.get('options') and any(q['options'].values()):
                question.set_options(q['options'])
            
            db.session.add(question)
            question_number += 1
            success_count += 1
            questions_data.append({
                'question_type': q['question_type'],
                'content': q['content'],
                'score': q['score'],
                'correct_answer': q.get('correct_answer', ''),
                'options': q.get('options', {'A': '', 'B': '', 'C': '', 'D': ''})
            })
        
        db.session.commit()
        
        return jsonify({
            'success': True,
            'message': f'成功导入 {success_count} 道题目',
            'success_count': success_count,
            'error_count': 0,
            'questions': questions_data
        }), 200
        
    except Exception as e:
        db.session.rollback()
        return jsonify({'error': f'导入失败：{str(e)}'}), 500
