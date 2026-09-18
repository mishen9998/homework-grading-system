"""Organization-scoped personnel and class administration."""
from functools import wraps
from io import BytesIO
import secrets
import string

import openpyxl
from flask import Blueprint, jsonify, request, send_file
from sqlalchemy.exc import IntegrityError, SQLAlchemyError
from werkzeug.security import generate_password_hash

from app import db
from app.models import Course, Organization, OrganizationClass, Schedule, User
from app.services.organization_context import (
    PolicyError, audit, current_organization, lock_current_identity, scoped_query,
)
from app.utils import role_required

bp = Blueprint('admin', __name__, url_prefix='/api/admin')
admin_required = role_required('admin')
USER_ROLES = {'student', 'teacher', 'admin'}
TEXT_FIELDS = {
    'username': 80, 'email': 120, 'name': 100, 'student_id': 20,
    'teacher_id': 20, 'phone': 20, 'qq': 20, 'college': 100,
}
USER_FIELDS = set(TEXT_FIELDS) | {
    'password', 'role', 'is_active', 'organization_class_id', 'class_name',
}


def invalid(message, status=400, code='invalid_request'):
    raise PolicyError(message, status, code)


def object_body(allowed):
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        invalid('请求必须为 JSON 对象')
    reject_unknown(data, allowed)
    return data


def reject_unknown(data, allowed):
    unknown = set(data) - set(allowed)
    if unknown:
        invalid('不允许修改字段：' + ', '.join(sorted(unknown)))


def clean_text(value, name, limit, *, required=False):
    if value is None and not required:
        return None
    if not isinstance(value, str):
        invalid(f'{name} 必须为字符串')
    value = value.strip()
    if required and not value:
        invalid(f'{name} 不能为空')
    if len(value) > limit:
        invalid(f'{name} 不能超过 {limit} 个字符')
    return value or None


def validate_password(password):
    if (not isinstance(password, str) or not 12 <= len(password) <= 256 or
            len(set(password)) < 4 or sum((any(c.islower() for c in password),
                any(c.isupper() for c in password), any(c.isdigit() for c in password),
                any(not c.isalnum() and not c.isspace() for c in password))) < 3):
        invalid('密码需为 12–256 个字符，包含大小写字母、数字、符号中的至少 3 类')
    return password


def transactional(fn):
    @wraps(fn)
    def wrapped(*args, **kwargs):
        try:
            result = fn(*args, **kwargs)
            db.session.commit()
            return result
        except IntegrityError:
            db.session.rollback()
            return jsonify({'error': '数据冲突或仍有关联记录，请检查重复账号、邮箱或先停用账号',
                            'code': 'data_conflict'}), 409
        except SQLAlchemyError:
            db.session.rollback()
            return jsonify({'error': '保存失败，请稍后重试', 'code': 'database_error'}), 500
        except Exception:
            db.session.rollback()
            raise
    return wrapped


def lock_organization(organization_id=None):
    own_organization = organization_id is None
    organization_id = organization_id or current_organization().id
    organization = Organization.query.filter_by(id=organization_id).populate_existing().with_for_update().first()
    if organization is None:
        invalid('机构不存在', 404, 'not_found')
    if own_organization:
        lock_current_identity(roles=('admin',), write=True)
    return organization


def organization_users(organization):
    return User.query.filter_by(organization_id=organization.id)


def current_user_count(organization):
    # A locking read is essential on MySQL REPEATABLE READ: authentication may
    # already have opened an older snapshot before the organization lock.
    return len(organization_users(organization).with_entities(User.id).with_for_update().all())


def ensure_user_capacity(organization, additional=1):
    if current_user_count(organization) + additional > organization.user_limit:
        invalid('机构账号人数额度不足（停用账号仍计入总人数）', 409, 'user_limit_exceeded')


def user_values(data, organization, *, existing=None):
    reject_unknown(data, USER_FIELDS)
    if existing is None:
        for key in ('username', 'email', 'password', 'role', 'name'):
            if key not in data:
                invalid(f'缺少字段：{key}')
    values = {}
    for field, limit in TEXT_FIELDS.items():
        if field in data:
            values[field] = clean_text(data[field], field, limit,
                                       required=field in ('username', 'email', 'name'))
    if 'email' in values and ('@' not in values['email'] or any(c.isspace() for c in values['email'])):
        invalid('邮箱格式无效')
    if 'role' in data:
        if not isinstance(data['role'], str) or data['role'] not in USER_ROLES:
            invalid('角色仅允许 student、teacher、admin')
        values['role'] = data['role']
    if 'is_active' in data:
        if not isinstance(data['is_active'], bool):
            invalid('is_active 必须为布尔值')
        values['is_active'] = data['is_active']
    if 'password' in data:
        values['password'] = generate_password_hash(validate_password(data['password']))
    for field in ('username', 'email', 'student_id', 'teacher_id'):
        if values.get(field):
            query = organization_users(organization).filter(getattr(User, field) == values[field])
            if existing is not None:
                query = query.filter(User.id != existing.id)
            if query.with_for_update().first():
                invalid(f'本机构 {field} 已存在', 409, 'duplicate_user')
    if 'organization_class_id' in data or 'class_name' in data:
        class_id = data.get('organization_class_id')
        class_name = clean_text(data.get('class_name'), 'class_name', 100)
        classroom = None
        if 'organization_class_id' in data and class_id is not None:
            if type(class_id) is not int or class_id <= 0:
                invalid('organization_class_id 必须为正整数或 null')
            classroom = OrganizationClass.query.filter_by(id=class_id, organization_id=organization.id).populate_existing().with_for_update().first()
            if classroom is None:
                invalid('班级不存在或不属于本机构', 400, 'invalid_class')
        elif class_name:
            classroom = OrganizationClass.query.filter_by(name=class_name, organization_id=organization.id).populate_existing().with_for_update().first()
            if classroom is None:
                invalid('请先创建本机构班级', 400, 'invalid_class')
        if classroom and class_name and classroom.name != class_name:
            invalid('班级 ID 与名称不一致')
        if 'organization_class_id' in data and class_id is None and class_name:
            invalid('清空班级 ID 时不能同时指定班级名称')
        values['organization_class_id'] = classroom.id if classroom else None
        values['class_name'] = classroom.name if classroom else None
    return values


def new_user(values, organization):
    user = User(organization_id=organization.id, token_version=1,
                friend_code=''.join(secrets.choice(string.ascii_uppercase + string.digits) for _ in range(10)),
                **values)
    db.session.add(user)
    return user


def scoped_user(user_id, *, lock=False):
    query = scoped_query(User).filter_by(id=user_id)
    if lock:
        query = query.populate_existing().with_for_update()
    user = query.first()
    if user is None or user.role == 'platform_admin':
        invalid('用户不存在', 404, 'not_found')
    return user


def protect_last_admin(user, organization, *, new_role=None, new_active=None, deleting=False):
    removing = deleting or new_role not in (None, 'admin') or new_active is False
    if user.role == 'admin' and user.is_active and removing:
        others = organization_users(organization).with_entities(User.id).filter(
            User.role == 'admin', User.is_active.is_(True), User.id != user.id).with_for_update().all()
        if not others:
            invalid('不能删除、停用或降级本机构最后一位活跃管理员', 409, 'last_active_admin')


def pagination(query):
    try:
        page = int(request.args.get('page', '1'))
        per_page = int(request.args.get('per_page', '20'))
    except ValueError:
        invalid('分页参数必须为整数')
    if page < 1 or not 1 <= per_page <= 100:
        invalid('page 必须大于 0，per_page 范围为 1–100')
    return query.paginate(page=page, per_page=per_page, error_out=False)


@bp.get('/users')
@admin_required
def get_users():
    query = scoped_query(User).filter(User.role.in_(USER_ROLES))
    role = request.args.get('role')
    if role:
        if role not in USER_ROLES:
            invalid('角色筛选无效')
        query = query.filter_by(role=role)
    search = request.args.get('search', '').strip()
    if search:
        query = query.filter(db.or_(User.username.contains(search), User.name.contains(search),
                                   User.email.contains(search), User.student_id.contains(search)))
    page = pagination(query.order_by(User.id.desc()))
    return jsonify({'users': [user.to_dict() for user in page.items], 'total': page.total,
                    'page': page.page, 'per_page': page.per_page, 'has_next': page.has_next})


@bp.post('/users')
@admin_required
@transactional
def create_user():
    organization = lock_organization()
    data = object_body(USER_FIELDS)
    values = user_values(data, organization)
    ensure_user_capacity(organization)
    user = new_user(values, organization)
    db.session.flush()
    audit('user.created', target_type='user', target_id=user.id, details={'role': user.role})
    return jsonify({'message': '创建成功', 'user': user.to_dict()}), 201


@bp.get('/users/<int:user_id>')
@admin_required
def get_user(user_id):
    return jsonify({'user': scoped_user(user_id).to_dict()})


@bp.put('/users/<int:user_id>')
@admin_required
@transactional
def update_user(user_id):
    organization = lock_organization()
    user = scoped_user(user_id, lock=True)
    data = object_body(USER_FIELDS)
    values = user_values(data, organization, existing=user)
    protect_last_admin(user, organization, new_role=values.get('role'), new_active=values.get('is_active'))
    invalidate = ('password' in values or
                  ('role' in values and values['role'] != user.role) or
                  ('is_active' in values and values['is_active'] != user.is_active))
    for key, value in values.items():
        setattr(user, key, value)
    if invalidate:
        user.token_version += 1
    audit('user.updated', target_type='user', target_id=user.id,
          details={'fields': sorted(data), 'sessions_revoked': invalidate})
    return jsonify({'message': '更新成功', 'user': user.to_dict()})


@bp.delete('/users/<int:user_id>')
@admin_required
@transactional
def delete_user(user_id):
    organization = lock_organization()
    user = scoped_user(user_id, lock=True)
    protect_last_admin(user, organization, deleting=True)
    if user.is_active:
        user.is_active = False
        user.token_version += 1
    audit('user.deactivated', target_type='user', target_id=user.id,
          details={'via': 'delete', 'history_retained': True})
    return jsonify({'message': '账号已停用，历史记录保留', 'user': user.to_dict()})


IMPORT_COLUMNS = {
    '账号': 'username', '密码': 'password', '邮箱': 'email', '角色': 'role',
    '姓名': 'name', '学号': 'student_id', '工号': 'teacher_id', '电话': 'phone',
    'QQ': 'qq', '班级': 'class_name', '学院': 'college',
}


@bp.post('/import')
@admin_required
@transactional
def import_users():
    upload = request.files.get('file')
    if upload is None or not upload.filename.lower().endswith('.xlsx'):
        invalid('请上传 .xlsx 文件')
    try:
        workbook = openpyxl.load_workbook(BytesIO(upload.read()), read_only=True, data_only=False)
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        headers = list(next(rows))
        if len(headers) != len(set(headers)) or any(header not in IMPORT_COLUMNS for header in headers):
            invalid('导入列重复或包含不支持的字段')
        if not {'账号', '密码', '邮箱', '角色', '姓名'}.issubset(headers):
            invalid('Excel 缺少账号、密码、邮箱、角色或姓名列')
        payloads = []
        for index, row in enumerate(rows, 2):
            if not any(value is not None for value in row):
                continue
            if len(payloads) >= 1000:
                invalid('每次最多导入 1000 个账号')
            if any(isinstance(value, str) and value.startswith('=') for value in row):
                invalid(f'第 {index} 行不允许公式')
            payloads.append((index, {IMPORT_COLUMNS[key]: str(value).strip() if value is not None else None
                                     for key, value in zip(headers, row)}))
        workbook.close()
    except PolicyError:
        raise
    except Exception:
        invalid('Excel 文件无效或无法解析')
    if not payloads:
        invalid('没有可导入的账号')
    organization = lock_organization()
    ensure_user_capacity(organization, len(payloads))
    for index, payload in payloads:
        try:
            values = user_values(payload, organization)
            new_user(values, organization)
            db.session.flush()
        except PolicyError as error:
            invalid(f'第 {index} 行：{error.description}；本次导入未保存任何账号',
                    error.code, error.error_code)
    audit('users.imported', target_type='organization', target_id=organization.id,
          details={'count': len(payloads)})
    return jsonify({'message': f'成功导入 {len(payloads)} 个用户', 'success_count': len(payloads),
                    'error_count': 0, 'errors': []})


@bp.get('/template')
@admin_required
def download_template():
    workbook = openpyxl.Workbook()
    workbook.active.title = '用户导入模板'
    workbook.active.append(list(IMPORT_COLUMNS))
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(output, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name='用户导入模板.xlsx')


@bp.get('/statistics')
@admin_required
def get_statistics():
    counts = dict(scoped_query(User).with_entities(User.role, db.func.count(User.id)).group_by(User.role).all())
    return jsonify({'total_users': sum(counts.values()), 'student_count': counts.get('student', 0),
                    'teacher_count': counts.get('teacher', 0), 'admin_count': counts.get('admin', 0)})


def class_dict(classroom):
    return {'id': classroom.id, 'organization_id': classroom.organization_id, 'name': classroom.name}


def scoped_class(class_id, *, lock=False):
    query = scoped_query(OrganizationClass).filter_by(id=class_id)
    if lock:
        query = query.populate_existing().with_for_update()
    classroom = query.first()
    if classroom is None:
        invalid('班级不存在', 404, 'not_found')
    return classroom


@bp.get('/classes')
@admin_required
def get_classes():
    page = pagination(scoped_query(OrganizationClass).order_by(OrganizationClass.id))
    return jsonify({'classes': [class_dict(item) for item in page.items], 'total': page.total,
                    'page': page.page, 'per_page': page.per_page, 'has_next': page.has_next})


@bp.get('/classes/<int:class_id>')
@admin_required
def get_class(class_id):
    return jsonify({'class': class_dict(scoped_class(class_id))})


@bp.post('/classes')
@admin_required
@transactional
def create_class():
    organization = lock_organization()
    data = object_body({'name'})
    name = clean_text(data.get('name'), 'name', 100, required=True)
    if scoped_query(OrganizationClass).filter_by(name=name).first():
        invalid('本机构班级名已存在', 409, 'duplicate_class')
    classroom = OrganizationClass(organization_id=organization.id, name=name)
    db.session.add(classroom)
    db.session.flush()
    audit('class.created', target_type='class', target_id=classroom.id)
    return jsonify({'class': class_dict(classroom)}), 201


@bp.put('/classes/<int:class_id>')
@admin_required
@transactional
def update_class(class_id):
    lock_organization()
    classroom = scoped_class(class_id, lock=True)
    data = object_body({'name'})
    name = clean_text(data.get('name'), 'name', 100, required=True)
    if scoped_query(OrganizationClass).filter(OrganizationClass.name == name,
                                             OrganizationClass.id != class_id).first():
        invalid('本机构班级名已存在', 409, 'duplicate_class')
    if len(name) > 50 and scoped_query(Schedule).filter_by(class_name=classroom.name).with_for_update().first():
        invalid('关联旧课表的班级名称不能超过 50 个字符')
    old_name = classroom.name
    classroom.name = name
    scoped_query(User).filter_by(organization_class_id=class_id).update({'class_name': name}, synchronize_session=False)
    # Course/schedule class names are legacy display/matching fields until their
    # own entity migrations; keep those names coherent inside this organization.
    for model in (Course, Schedule):
        scoped_query(model).filter_by(class_name=old_name).update({'class_name': name}, synchronize_session=False)
    audit('class.updated', target_type='class', target_id=class_id)
    return jsonify({'class': class_dict(classroom)})


@bp.delete('/classes/<int:class_id>')
@admin_required
@transactional
def delete_class(class_id):
    lock_organization()
    classroom = scoped_class(class_id, lock=True)
    if (scoped_query(User).filter_by(organization_class_id=class_id).first() or
            any(scoped_query(model).filter_by(class_name=classroom.name).first() for model in (Course, Schedule))):
        invalid('班级仍有关联人员、课程或课表，请先调整关联', 409, 'class_in_use')
    db.session.delete(classroom)
    audit('class.deleted', target_type='class', target_id=class_id)
    return jsonify({'message': '删除成功'})
