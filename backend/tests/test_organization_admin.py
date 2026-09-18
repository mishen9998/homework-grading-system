"""Synthetic administration contracts. No .env, service or provider access."""
from datetime import datetime, timedelta
from io import BytesIO
import os
import unittest
from unittest.mock import patch

import openpyxl
from flask_jwt_extended import create_access_token
from sqlalchemy.exc import IntegrityError
from werkzeug.security import check_password_hash

from app import create_app, db
from app.models import AuditLog, Course, Organization, OrganizationClass, Schedule, User
from app.services.organization_context import session_claims
from config import TestingConfig


class AdminTestConfig(TestingConfig):
    SQLALCHEMY_DATABASE_URI = 'sqlite:///:memory:'
    SQLALCHEMY_ENGINE_OPTIONS = {}
    SECRET_KEY = 'synthetic-admin-test-secret-32-characters'
    JWT_SECRET_KEY = 'synthetic-admin-jwt-secret-32-characters'
    AUTO_CREATE_TABLES = True
    REDIS_URL = ''


class OrganizationAdminTests(unittest.TestCase):
    def setUp(self):
        with patch('app._setup_logging'):
            self.app = create_app(AdminTestConfig)
        self.context = self.app.app_context()
        self.context.push()
        self.client = self.app.test_client()
        self.org_a = Organization(code='alpha', name='Alpha', user_limit=10)
        self.org_b = Organization(code='beta', name='Beta', user_limit=10)
        self.platform_org = Organization(code='platform', name='Platform')
        db.session.add_all([self.org_a, self.org_b, self.platform_org])
        db.session.flush()
        self.admin_a = self.seed_user(self.org_a, 'admin', 'admin')
        self.admin_b = self.seed_user(self.org_b, 'admin', 'admin')
        self.student_a = self.seed_user(self.org_a, 'student', 'student')
        self.platform = self.seed_user(self.platform_org, 'platform', 'platform_admin')
        self.class_b = OrganizationClass(organization_id=self.org_b.id, name='Beta Class')
        db.session.add(self.class_b)
        db.session.commit()
        self.tokens = {key: self.token(user) for key, user in (
            ('a', self.admin_a), ('b', self.admin_b), ('s', self.student_a), ('p', self.platform))}

    def tearDown(self):
        db.session.remove()
        db.drop_all()
        db.engine.dispose()
        self.context.pop()

    def seed_user(self, organization, username, role):
        user = User(organization_id=organization.id, username=username,
                    email=username + '@example.test', name=username, role=role,
                    password='synthetic-unusable-hash', is_active=True, token_version=1)
        db.session.add(user)
        db.session.flush()
        return user

    def token(self, user):
        return create_access_token(identity=str(user.id), additional_claims=session_claims(user))

    def request(self, method, path, *, who='a', **kwargs):
        return self.client.open(path, method=method,
                                headers={'Authorization': 'Bearer ' + self.tokens[who]}, **kwargs)

    def payload(self, username='new-user', **extra):
        return {'username': username, 'email': username + '@example.test', 'name': 'Synthetic User',
                'password': 'Synthetic!Password42', 'role': 'student', **extra}

    def organization_payload(self, **extra):
        admin = self.payload('first-admin')
        admin.pop('role')
        return {'code': 'customer', 'name': 'Customer', 'expires_at': '2030-01-01T08:00:00+08:00',
                'user_limit': 5, 'admin': admin, **extra}

    def test_scoped_lists_counts_and_cross_org_objects(self):
        result = self.request('GET', '/api/admin/users')
        self.assertEqual(result.status_code, 200)
        self.assertEqual(result.json['total'], 2)
        self.assertTrue(all(row['organization_id'] == self.org_a.id for row in result.json['users']))
        self.assertEqual(self.request('GET', '/api/admin/statistics').json['total_users'], 2)
        for method in ('GET', 'PUT', 'DELETE'):
            with self.subTest(method=method):
                response = self.request(method, f'/api/admin/users/{self.admin_b.id}', json={'name': 'Changed'})
                self.assertEqual(response.status_code, 404)

    def test_roles_and_platform_boundary(self):
        self.assertEqual(self.request('GET', '/api/admin/users', who='s').status_code, 403)
        self.assertEqual(self.request('GET', '/api/admin/users', who='p').status_code, 403)
        self.assertEqual(self.request('GET', '/api/platform/organizations', who='a').status_code, 403)
        response = self.request('GET', '/api/platform/organizations', who='p')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json['total'], 2)
        self.assertNotIn('users', response.json)

    def test_create_scopes_identity_and_rejects_authority_fields(self):
        for extra in ({'organization_id': self.org_b.id}, {'role': 'platform_admin'}, {'token_version': 99},
                      {'is_active': 'true'}, {'role': []}, {'password': '123456789012'}):
            with self.subTest(extra=extra):
                response = self.request('POST', '/api/admin/users', json=self.payload(**extra))
                self.assertEqual(response.status_code, 400)
        response = self.request('POST', '/api/admin/users', json=self.payload())
        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.json['user']['organization_id'], self.org_a.id)
        stored = db.session.get(User, response.json['user']['id'])
        self.assertTrue(check_password_hash(stored.password, 'Synthetic!Password42'))
        self.assertNotIn('password', response.json['user'])

    def test_username_email_uniqueness_is_per_org(self):
        response = self.request('POST', '/api/admin/users', json=self.payload('same'))
        self.assertEqual(response.status_code, 201)
        self.assertEqual(self.request('POST', '/api/admin/users', who='b', json=self.payload('same')).status_code, 201)
        self.assertEqual(self.request('POST', '/api/admin/users', json=self.payload('same')).status_code, 409)

    def test_user_quota_includes_inactive_accounts(self):
        self.org_a.user_limit = 2
        self.student_a.is_active = False
        db.session.commit()
        response = self.request('POST', '/api/admin/users', json=self.payload())
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.json['code'], 'user_limit_exceeded')
        self.assertEqual(User.query.filter_by(organization_id=self.org_a.id).count(), 2)

    def test_last_active_admin_cannot_be_removed(self):
        for method, data in (('DELETE', {}), ('PUT', {'is_active': False}), ('PUT', {'role': 'teacher'})):
            with self.subTest(method=method, data=data):
                response = self.request(method, f'/api/admin/users/{self.admin_a.id}', json=data)
                self.assertEqual(response.status_code, 409)
                self.assertEqual(response.json['code'], 'last_active_admin')
        self.seed_user(self.org_a, 'inactive-admin', 'admin').is_active = False
        db.session.commit()
        self.assertEqual(self.request('DELETE', f'/api/admin/users/{self.admin_a.id}').status_code, 409)

    def test_password_role_and_inactive_changes_revoke_tokens(self):
        for change in ({'password': 'Replacement!Password42'}, {'role': 'teacher'}, {'is_active': False}):
            with self.subTest(change=change):
                old_version = self.student_a.token_version
                self.tokens['s'] = self.token(self.student_a)
                response = self.request('PUT', f'/api/admin/users/{self.student_a.id}', json=change)
                self.assertEqual(response.status_code, 200)
                self.assertEqual(self.student_a.token_version, old_version + 1)
                self.assertEqual(self.request('GET', '/api/admin/organization', who='s').status_code, 401)

    def test_delete_deactivates_and_retains_history(self):
        response = self.request('DELETE', f'/api/admin/users/{self.student_a.id}')
        self.assertEqual(response.status_code, 200)
        self.assertIsNotNone(db.session.get(User, self.student_a.id))
        self.assertFalse(self.student_a.is_active)
        self.assertEqual(self.request('GET', '/api/admin/organization', who='s').status_code, 401)

    def test_classes_crud_assignment_and_cross_org_rejection(self):
        response = self.request('POST', '/api/admin/classes', json={'name': 'Class One'})
        self.assertEqual(response.status_code, 201)
        class_id = response.json['class']['id']
        response = self.request('PUT', f'/api/admin/users/{self.student_a.id}',
                                json={'organization_class_id': self.class_b.id})
        self.assertEqual(response.status_code, 400)
        self.assertIsNone(self.student_a.organization_class_id)
        response = self.request('PUT', f'/api/admin/users/{self.student_a.id}',
                                json={'organization_class_id': class_id})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json['user']['class_name'], 'Class One')
        self.assertEqual(self.request('DELETE', f'/api/admin/classes/{class_id}').status_code, 409)
        self.assertEqual(self.request('PUT', f'/api/admin/classes/{class_id}', json={'name': 'Renamed'}).status_code, 200)
        db.session.expire_all()
        self.assertEqual(self.student_a.class_name, 'Renamed')
        for method in ('GET', 'PUT', 'DELETE'):
            self.assertEqual(self.request(method, f'/api/admin/classes/{self.class_b.id}', json={'name': 'Wrong'}).status_code, 404)
        self.assertEqual(self.request('PUT', f'/api/admin/users/{self.student_a.id}',
                                      json={'organization_class_id': None}).status_code, 200)
        self.assertEqual(self.request('DELETE', f'/api/admin/classes/{class_id}').status_code, 200)

    def test_class_rename_keeps_legacy_course_and_schedule_names_scoped(self):
        classroom = OrganizationClass(organization_id=self.org_a.id, name='Shared Name')
        db.session.add(classroom)
        course = Course(organization_id=self.org_a.id, teacher_id=self.admin_a.id,
                        name='Synthetic course', code='CODE-A', class_name='Shared Name')
        schedule_a = Schedule(organization_id=self.org_a.id, course_name='Synthetic',
                              week_day=1, period=1, class_name='Shared Name')
        schedule_b = Schedule(organization_id=self.org_b.id, course_name='Synthetic',
                              week_day=1, period=1, class_name='Shared Name')
        db.session.add_all([course, schedule_a, schedule_b])
        db.session.commit()
        self.assertEqual(self.request('DELETE', f'/api/admin/classes/{classroom.id}').status_code, 409)
        too_long = self.request('PUT', f'/api/admin/classes/{classroom.id}', json={'name': 'X' * 51})
        self.assertEqual(too_long.status_code, 400)
        self.assertEqual(classroom.name, 'Shared Name')
        self.assertEqual(schedule_a.class_name, 'Shared Name')
        response = self.request('PUT', f'/api/admin/classes/{classroom.id}', json={'name': 'Updated'})
        self.assertEqual(response.status_code, 200)
        db.session.expire_all()
        self.assertEqual(course.class_name, 'Updated')
        self.assertEqual(schedule_a.class_name, 'Updated')
        self.assertEqual(schedule_b.class_name, 'Shared Name')

    def workbook(self, rows):
        workbook = openpyxl.Workbook()
        workbook.active.append(['账号', '密码', '邮箱', '角色', '姓名'])
        for row in rows:
            workbook.active.append(row)
        stream = BytesIO()
        workbook.save(stream)
        stream.seek(0)
        return stream

    def import_rows(self, rows):
        return self.request('POST', '/api/admin/import', data={'file': (self.workbook(rows), 'users.xlsx')})

    def import_row(self, username='import-one', role='student'):
        return [username, 'Synthetic!Password42', username + '@example.test', role, 'Synthetic']

    def test_import_is_atomic_on_invalid_row_or_duplicate(self):
        for rows in ([self.import_row(), self.import_row('bad', 'platform_admin')],
                     [self.import_row(), self.import_row()],
                     [self.import_row(), ['=formula', 'Synthetic!Password42', 'f@example.test', 'student', 'Name']]):
            response = self.import_rows(rows)
            self.assertIn(response.status_code, (400, 409))
            self.assertEqual(User.query.filter_by(organization_id=self.org_a.id).count(), 2)
            self.assertEqual(AuditLog.query.filter_by(action='users.imported').count(), 0)

    def test_import_success_and_quota_rollback(self):
        self.org_a.user_limit = 3
        db.session.commit()
        self.assertEqual(self.import_rows([self.import_row(), self.import_row('import-two')]).status_code, 409)
        self.assertEqual(User.query.filter_by(organization_id=self.org_a.id).count(), 2)
        response = self.import_rows([self.import_row()])
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json['success_count'], 1)
        self.assertEqual(AuditLog.query.filter_by(action='users.imported').count(), 1)

    def test_template_contains_headers_without_shared_default_password(self):
        response = self.request('GET', '/api/admin/template')
        self.assertEqual(response.status_code, 200)
        workbook = openpyxl.load_workbook(BytesIO(response.data))
        self.assertEqual(workbook.active.max_row, 1)
        self.assertIn('工号', [cell.value for cell in workbook.active[1]])

    def test_platform_create_first_admin_defaults_and_audit(self):
        response = self.request('POST', '/api/platform/organizations', who='p', json=self.organization_payload())
        self.assertEqual(response.status_code, 201, response.json)
        organization = response.json['organization']
        self.assertFalse(organization['ai_enabled'])
        self.assertEqual(organization['user_count'], 1)
        self.assertEqual(organization['expires_at'], '2030-01-01T00:00:00')
        user = db.session.get(User, response.json['admin']['id'])
        self.assertEqual(user.role, 'admin')
        self.assertEqual(user.organization_id, organization['id'])
        audit = AuditLog.query.filter_by(action='organization.created').one()
        self.assertEqual(audit.actor_id, self.platform.id)
        self.assertNotIn('password', str(audit.details))

    def test_platform_creation_rolls_back_invalid_admin_or_zero_quota(self):
        for extra in ({'user_limit': 0}, {'admin': {'username': 'missing'}}, {'ai_enabled': True},
                      {'code': 'platform'}, {'user_limit': True}):
            response = self.request('POST', '/api/platform/organizations', who='p',
                                    json=self.organization_payload(**extra))
            self.assertIn(response.status_code, (400, 409), response.json)
            self.assertEqual(Organization.query.count(), 3)
            self.assertEqual(User.query.count(), 4)

    def test_platform_update_customer_lifecycle_and_quotas(self):
        response = self.request('PUT', f'/api/platform/organizations/{self.org_a.id}', who='p',
                                json={'status': 'disabled', 'user_limit': 6, 'storage_limit_bytes': 5000,
                                      'ai_monthly_token_limit': 1234, 'concurrent_task_limit': 1,
                                      'expires_at': '2031-01-01T00:00:00Z'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.request('GET', '/api/admin/users').status_code, 403)
        self.assertEqual(self.request('GET', f'/api/platform/organizations/{self.org_a.id}', who='p').status_code, 200)
        self.assertEqual(self.request('PUT', f'/api/platform/organizations/{self.org_a.id}', who='p',
                                      json={'status': 'active', 'expires_at': None}).status_code, 200)
        for payload in ({'user_limit': 1}, {'ai_enabled': True}, {'code': 'changed'},
                        {'storage_limit_bytes': -1}, {'concurrent_task_limit': 0.5},
                        {'expires_at': '0001-01-01T00:00:00Z'},
                        {'expires_at': '0001-01-01T00:00:00+14:00'},
                        {'expires_at': '1000-01-01T00:00:00+14:00'}):
            self.assertIn(self.request('PUT', f'/api/platform/organizations/{self.org_a.id}', who='p',
                                       json=payload).status_code, (400, 409))
        self.assertEqual(self.request('PUT', f'/api/platform/organizations/{self.platform_org.id}',
                                      who='p', json={'status': 'disabled'}).status_code, 404)

    def test_ai_authorization_only_own_admin_can_change(self):
        self.assertEqual(self.request('GET', '/api/admin/organization', who='s').status_code, 200)
        self.assertEqual(self.request('PUT', '/api/admin/organization', who='s', json={'ai_enabled': True}).status_code, 403)
        for body in ({'ai_enabled': 'true'}, {'ai_enabled': True, 'organization_id': self.org_b.id},
                     {'ai_enabled': True, 'user_limit': 100}):
            self.assertEqual(self.request('PUT', '/api/admin/organization', json=body).status_code, 400)
        self.assertEqual(self.request('PUT', '/api/admin/organization', json={'ai_enabled': True}).status_code, 200)
        self.assertTrue(self.org_a.ai_enabled)
        self.assertFalse(self.org_b.ai_enabled)
        self.assertEqual(AuditLog.query.filter_by(action='organization.ai_authorization_changed').count(), 1)

    def test_expiration_read_only_and_disabled_denial(self):
        self.org_a.expires_at = datetime.utcnow() - timedelta(days=1)
        db.session.commit()
        self.assertEqual(self.request('GET', '/api/admin/users').status_code, 200)
        self.assertEqual(self.request('GET', '/api/admin/template').status_code, 200)
        self.assertEqual(self.request('POST', '/api/admin/users', json=self.payload()).status_code, 403)
        self.assertEqual(self.request('PUT', '/api/admin/organization', json={'ai_enabled': True}).status_code, 403)
        self.org_a.status = 'disabled'
        db.session.commit()
        self.assertEqual(self.request('GET', '/api/admin/users').status_code, 403)

    def test_database_failure_rolls_back_user_and_audit(self):
        with patch.object(db.session, 'commit', side_effect=IntegrityError('synthetic', {}, Exception('conflict'))):
            response = self.request('POST', '/api/admin/users', json=self.payload())
        self.assertEqual(response.status_code, 409)
        self.assertIsNone(User.query.filter_by(username='new-user').first())
        self.assertEqual(AuditLog.query.filter_by(action='user.created').count(), 0)

    def test_bootstrap_independent_account_preserves_old_admin(self):
        from create_admin import create_platform_admin
        user_id = create_platform_admin('admin', 'Bootstrap!Password42', 'root@example.test')
        created = db.session.get(User, user_id)
        self.assertEqual(created.role, 'platform_admin')
        self.assertEqual(created.organization_id, self.platform_org.id)
        self.assertEqual(self.admin_a.role, 'admin')
        self.assertEqual(self.admin_a.organization_id, self.org_a.id)
        with self.assertRaises(Exception):
            create_platform_admin('admin', 'Bootstrap!Password42', 'root@example.test')

    def test_bootstrap_requires_explicit_database_before_app_creation(self):
        from create_admin import main
        with patch.dict(os.environ, {}, clear=True), patch('create_admin.create_app') as factory:
            with self.assertRaises(SystemExit):
                main()
            factory.assert_not_called()

    def test_invalid_pagination_and_json_are_rejected(self):
        for query in ('page=0', 'page=no', 'per_page=101', 'per_page=-1'):
            self.assertEqual(self.request('GET', '/api/admin/users?' + query).status_code, 400)
        self.assertEqual(self.request('POST', '/api/admin/users', json=[]).status_code, 400)


if __name__ == '__main__':
    unittest.main()
